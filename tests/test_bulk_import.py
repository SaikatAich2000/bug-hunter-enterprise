"""Tests for the bulk-import feature (app/bulk_import.py + the two routes on
app/routes/bugs.py: GET /api/bugs/import/template.xlsx, POST /api/bugs/import).
"""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from app.bulk_import import TEMPLATE_HEADERS
from tests.conftest import BOOTSTRAP_EMAIL

_TEMPLATE = "/api/bugs/import/template.xlsx"
_IMPORT = "/api/bugs/import"


def _make_project(client, name="Import Proj"):
    r = client.post("/api/projects", json={"name": name, "color": "#000000"})
    assert r.status_code == 201, r.text
    return r.json()


def _make_user(client, name, email, role="user", password="UserPass99"):
    r = client.post("/api/users", json={
        "name": name, "email": email, "role": role, "password": password,
    })
    assert r.status_code == 201, r.text
    return r.json()


def _xlsx_bytes(rows: list[dict[str, str]]) -> bytes:
    """Build an upload matching the exact template layout: banner + header
    row + one row per dict (missing keys become blank cells)."""
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Item Detail Export — banner row")
    for col, label in enumerate(TEMPLATE_HEADERS, start=1):
        ws.cell(row=2, column=col, value=label)
    for r_idx, row in enumerate(rows, start=3):
        for col, label in enumerate(TEMPLATE_HEADERS, start=1):
            ws.cell(row=r_idx, column=col, value=row.get(label, ""))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_bytes(rows: list[dict[str, str]]) -> bytes:
    lines = [",".join(f'"{h}"' for h in TEMPLATE_HEADERS)]
    for row in rows:
        lines.append(",".join(f'"{row.get(h, "")}"' for h in TEMPLATE_HEADERS))
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


def _upload(client, data: bytes, filename: str):
    return client.post(
        _IMPORT,
        files={"file": (filename, io.BytesIO(data), "application/octet-stream")},
    )


# --- Auth / permission gates ---

def test_template_requires_auth(client):
    assert client.get(_TEMPLATE).status_code == 401


def test_import_requires_auth(client):
    assert _upload(client, _xlsx_bytes([]), "x.xlsx").status_code == 401


def test_template_forbidden_for_regular_user(user_client):
    assert user_client.get(_TEMPLATE).status_code == 403


def test_import_forbidden_for_regular_user(user_client):
    assert _upload(user_client, _xlsx_bytes([]), "x.xlsx").status_code == 403


# --- Template download ---

def test_download_template_matches_headers(admin_client):
    res = admin_client.get(_TEMPLATE)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.worksheets[0]
    header_row = [c.value for c in ws[2]]
    assert header_row == TEMPLATE_HEADERS
    # Zero data rows beyond the banner + header.
    assert ws.max_row == 2
    assert "Instructions" in wb.sheetnames


def test_download_template_build_failure_returns_500(admin_client, monkeypatch):
    import app.routes.bugs as bugs_mod
    def _boom():
        raise bugs_mod.bulk_import.BulkImportError("workbook build failed")
    monkeypatch.setattr(bugs_mod.bulk_import, "build_template_workbook", _boom)
    res = admin_client.get(_TEMPLATE)
    assert res.status_code == 500
    assert "workbook build failed" in res.json()["detail"]


# --- Happy path ---

def test_import_creates_bugs_xlsx(admin_client):
    proj = _make_project(admin_client, "Import Proj A")
    data = _xlsx_bytes([
        {"Title": "Bulk row one", "Project": proj["name"], "Priority": "High"},
        {
            "Title": "Bulk row two — a task", "Type": "Task",
            "Project": proj["name"], "Status": "In Progress",
        },
    ])
    res = _upload(admin_client, data, "bugs.xlsx")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 2
    assert body["failed"] == 0
    assert len(body["created_ids"]) == 2

    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    titles = {b["title"] for b in listed["items"]}
    assert "Bulk row one" in titles
    assert "Bulk row two — a task" in titles
    task = next(b for b in listed["items"] if b["item_type"] == "Task")
    assert task["status"] == "In Progress"
    bug = next(b for b in listed["items"] if b["item_type"] == "Bug")
    assert bug["priority"] == "High"
    # Default reporter is the importer when Reporter/Reporter Email are blank.
    assert bug["reporter"]["email"] == BOOTSTRAP_EMAIL


def test_import_creates_bugs_csv(admin_client):
    proj = _make_project(admin_client, "Import Proj CSV")
    data = _csv_bytes([{"Title": "CSV row", "Project": proj["name"]}])
    res = _upload(admin_client, data, "bugs.csv")
    assert res.status_code == 200, res.text
    assert res.json()["created"] == 1


def test_import_resolves_reporter_and_assignees_by_email(admin_client):
    proj = _make_project(admin_client, "Import Proj B")
    dev = _make_user(admin_client, "Dana Dev", "dana@test.local")
    data = _xlsx_bytes([{
        "Title": "Assigned row", "Project": proj["name"],
        "Reporter Email": dev["email"], "Assignees": dev["email"],
    }])
    res = _upload(admin_client, data, "bugs.xlsx")
    assert res.status_code == 200, res.text
    assert res.json()["created"] == 1
    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    row = listed["items"][0]
    assert row["reporter"]["email"] == dev["email"]
    assert [a["email"] for a in row["assignees"]] == [dev["email"]]


def test_import_unresolved_assignee_is_a_warning_not_a_failure(admin_client):
    proj = _make_project(admin_client, "Import Proj C")
    data = _xlsx_bytes([{
        "Title": "Row with bad assignee", "Project": proj["name"],
        "Assignees": "Nobody Here",
    }])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 0
    assert len(body["warnings"]) == 1
    assert "Nobody Here" in body["warnings"][0]["warning"]
    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    assert listed["items"][0]["assignees"] == []


# --- Row-level validation errors ---

def test_import_reports_row_errors_and_skips_them(admin_client):
    proj = _make_project(admin_client, "Import Proj D")
    data = _xlsx_bytes([
        {"Title": "ok row", "Project": proj["name"]},
        {"Title": "no such project", "Project": "Does Not Exist"},
        {"Title": "x", "Project": proj["name"]},  # title too short
        {"Title": "bad type row", "Project": proj["name"], "Type": "Epic"},
        {
            "Title": "bad status", "Project": proj["name"],
            "Status": "Not A Real Status",
        },
    ])
    res = _upload(admin_client, data, "bugs.xlsx")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 4
    errors_by_row = {e["row"]: e["error"] for e in body["errors"]}
    assert any("not found" in msg for msg in errors_by_row.values())
    assert any("at least" in msg for msg in errors_by_row.values())
    assert any("Invalid Type" in msg for msg in errors_by_row.values())
    assert any("Invalid Status" in msg for msg in errors_by_row.values())


# --- Every column is optional; blanks degrade gracefully, never crash ---

def test_import_blank_title_gets_a_placeholder(admin_client):
    proj = _make_project(admin_client, "Import Proj E")
    data = _xlsx_bytes([{"Title": "", "Project": proj["name"]}])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 0
    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    assert listed["items"][0]["title"]  # non-empty placeholder, not a crash


def test_import_blank_project_defaults_when_only_one_exists(admin_client):
    # A fresh app always seeds exactly one project ("General") on first boot
    # (see app/main.py::_bootstrap) — no need to create another to test this.
    projects = admin_client.get("/api/projects").json()
    assert len(projects) == 1
    general = projects[0]
    data = _xlsx_bytes([{"Title": "row with no project column value"}])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1, body
    listed = admin_client.get("/api/bugs", params={"project_id": general["id"]}).json()
    assert listed["total"] == 1


def test_import_blank_project_fails_cleanly_when_ambiguous(admin_client):
    # "General" (seeded on boot) plus one more makes the default ambiguous.
    _make_project(admin_client, "Proj One")
    data = _xlsx_bytes([{"Title": "ambiguous project row"}])
    res = _upload(admin_client, data, "bugs.xlsx")
    assert res.status_code == 200  # never a crash, just a reported row error
    body = res.json()
    assert body["created"] == 0
    assert body["failed"] == 1
    assert "Project is required" in body["errors"][0]["error"]


def test_import_every_optional_column_blank_still_creates(admin_client):
    proj = _make_project(admin_client, "Import Proj F")
    data = _xlsx_bytes([{"Title": "minimal row", "Project": proj["name"]}])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1, body
    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    row = listed["items"][0]
    assert row["item_type"] == "Bug"
    assert row["status"] == "New"
    assert row["priority"] == "Medium"
    assert row["environment"] == "DEV"
    assert row["due_date"] is None
    assert row["reporter"]["email"] == BOOTSTRAP_EMAIL
    assert row["assignees"] == []


# --- Whole-file problems ---

def test_import_rejects_unsupported_file_type(admin_client):
    res = _upload(admin_client, b"just some text", "bugs.txt")
    assert res.status_code == 400
    assert "Unsupported file type" in res.json()["detail"]


def test_import_rejects_missing_header_row(admin_client):
    bad_csv = b'"Foo","Bar"\r\n"1","2"\r\n'
    res = _upload(admin_client, bad_csv, "bugs.csv")
    assert res.status_code == 400
    assert "header row" in res.json()["detail"]


def test_import_empty_file_reports_no_rows(admin_client):
    data = _xlsx_bytes([])
    res = _upload(admin_client, data, "bugs.xlsx")
    assert res.status_code == 200
    body = res.json()
    assert body["created"] == 0
    assert body["failed"] == 0
    assert "No data rows" in body["message"]


# --- Row-level validation: priority / environment / due date ---

def test_import_invalid_priority_environment_due_date(admin_client):
    proj = _make_project(admin_client, "Import Proj Bad Enums")
    data = _xlsx_bytes([
        {"Title": "bad priority", "Project": proj["name"], "Priority": "Extreme"},
        {"Title": "bad environment", "Project": proj["name"], "Env": "MOON"},
        {"Title": "bad due date", "Project": proj["name"], "Due Date": "13/45/2026"},
        {"Title": "bad status", "Project": proj["name"], "Status": "Vaporized"},
        {"Title": "good due date", "Project": proj["name"], "Due Date": "2026-12-31"},
    ])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 4
    msgs = [e["error"] for e in body["errors"]]
    assert any("Invalid Priority" in m for m in msgs)
    assert any("Invalid Env" in m for m in msgs)
    assert any("YYYY-MM-DD" in m for m in msgs)
    assert any("Invalid Status" in m for m in msgs)

    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    good = next(b for b in listed["items"] if b["title"] == "good due date")
    assert good["due_date"] == "2026-12-31"


def test_import_deactivated_assignee_is_a_warning(admin_client):
    proj = _make_project(admin_client, "Import Proj Deactivated Assignee")
    inactive = _make_user(admin_client, "Inactive Assignee", "inactive-assign@test.local")
    admin_client.put(f"/api/users/{inactive['id']}", json={"is_active": False})
    data = _xlsx_bytes([{
        "Title": "deactivated assignee row", "Project": proj["name"],
        "Assignees": "Inactive Assignee",
    }])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 0
    assert "deactivated" in body["warnings"][0]["warning"]
    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    assert listed["items"][0]["assignees"] == []


# --- Row-level validation: reporter resolution by name ---

def test_import_reporter_by_name_resolved_not_found_ambiguous_deactivated(admin_client):
    proj = _make_project(admin_client, "Import Proj Reporter")
    solo = _make_user(admin_client, "Solo Reporter", "solo@test.local")
    _make_user(admin_client, "Dup Reporter", "dup1@test.local")
    _make_user(admin_client, "Dup Reporter", "dup2@test.local")
    inactive = _make_user(admin_client, "Inactive Reporter", "inactive-rep@test.local")
    admin_client.put(f"/api/users/{inactive['id']}", json={"is_active": False})

    data = _xlsx_bytes([
        {"Title": "resolved by name", "Project": proj["name"], "Reporter": "Solo Reporter"},
        {"Title": "reporter not found", "Project": proj["name"], "Reporter": "Ghost Person"},
        {"Title": "ambiguous reporter", "Project": proj["name"], "Reporter": "Dup Reporter"},
        {
            "Title": "deactivated reporter", "Project": proj["name"],
            "Reporter": "Inactive Reporter",
        },
        {
            "Title": "reporter email not found", "Project": proj["name"],
            "Reporter Email": "nobody@nowhere.test",
        },
    ])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 4
    msgs = [e["error"] for e in body["errors"]]
    assert any("Ghost Person" in m and "not found" in m for m in msgs)
    assert any("Dup Reporter" in m and "more than one" in m for m in msgs)
    assert any("Inactive Reporter" in m and "deactivated" in m for m in msgs)
    assert any("nobody@nowhere.test" in m and "not found" in m for m in msgs)

    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    resolved = next(b for b in listed["items"] if b["title"] == "resolved by name")
    assert resolved["reporter"]["email"] == solo["email"]


# --- Row-level validation: assignees ---

def test_import_assignees_ambiguous_blank_token_and_duplicate(admin_client):
    proj = _make_project(admin_client, "Import Proj Assignees")
    solo = _make_user(admin_client, "Solo Assignee", "solo-assignee@test.local")
    _make_user(admin_client, "Dup Assignee", "dupassign1@test.local")
    _make_user(admin_client, "Dup Assignee", "dupassign2@test.local")

    data = _xlsx_bytes([{
        "Title": "assignee edge cases", "Project": proj["name"],
        # blank token from the double comma, a valid single match, a
        # duplicate of that same match, and an ambiguous name.
        "Assignees": "Solo Assignee,,Solo Assignee,Dup Assignee",
    }])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 0
    assert len(body["warnings"]) == 1
    assert "Dup Assignee" in body["warnings"][0]["warning"]
    assert "more than one" in body["warnings"][0]["warning"]

    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    row = listed["items"][0]
    # Solo Assignee appears exactly once despite being listed twice.
    assert [a["email"] for a in row["assignees"]] == [solo["email"]]


# --- Row-level validation: event resolution ---

def test_import_event_resolved_not_found_and_ambiguous(admin_client):
    proj = _make_project(admin_client, "Import Proj Event")
    ev = admin_client.post("/api/events", json={"name": "Sprint Solo"}).json()
    admin_client.post("/api/events", json={"name": "Sprint Dup"})
    admin_client.post("/api/events", json={"name": "Sprint Dup"})

    data = _xlsx_bytes([
        {"Title": "event resolved", "Project": proj["name"], "Event": "Sprint Solo"},
        {"Title": "event not found", "Project": proj["name"], "Event": "No Such Sprint"},
        {"Title": "event ambiguous", "Project": proj["name"], "Event": "Sprint Dup"},
    ])
    res = _upload(admin_client, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 1
    assert body["failed"] == 2
    msgs = [e["error"] for e in body["errors"]]
    assert any("No Such Sprint" in m and "not found" in m for m in msgs)
    assert any("Sprint Dup" in m and "more than one" in m for m in msgs)

    listed = admin_client.get("/api/bugs", params={"project_id": proj["id"]}).json()
    resolved = next(b for b in listed["items"] if b["title"] == "event resolved")
    assert resolved["event_id"] == ev["id"]


# --- Low-level parsing helpers (app/bulk_import.py internals) ---

def test_ensure_openpyxl_raises_when_unavailable(monkeypatch):
    from app import bulk_import
    monkeypatch.setattr(bulk_import, "_OPENPYXL_AVAILABLE", False)
    try:
        bulk_import._ensure_openpyxl()
        raise AssertionError("expected BulkImportError")
    except bulk_import.BulkImportError as exc:
        assert "openpyxl is not installed" in str(exc)


def test_stringify_cell_formats_date_and_datetime():
    from datetime import date, datetime

    from app.bulk_import import _stringify_cell
    assert _stringify_cell(date(2026, 3, 4)) == "2026-03-04"
    assert _stringify_cell(datetime(2026, 3, 4, 9, 30)) == "2026-03-04"
    assert _stringify_cell(None) == ""
    assert _stringify_cell(" text  ") == "text"


def test_rows_from_xlsx_rejects_corrupt_file():
    from app.bulk_import import BulkImportError, _rows_from_xlsx
    try:
        _rows_from_xlsx(b"not a real xlsx file at all")
        raise AssertionError("expected BulkImportError")
    except BulkImportError as exc:
        assert "Excel workbook" in str(exc)


def test_rows_from_csv_rejects_invalid_utf8():
    from app.bulk_import import BulkImportError, _rows_from_csv
    try:
        _rows_from_csv(b"Title,Project\r\n\xff\xfebad,General\r\n")
        raise AssertionError("expected BulkImportError")
    except BulkImportError as exc:
        assert "UTF-8" in str(exc)


def test_rows_from_xlsx_stops_scanning_past_the_cap():
    from app.bulk_import import _MAX_SCAN_ROWS, _rows_from_xlsx
    wb = Workbook()
    ws = wb.active
    for i in range(_MAX_SCAN_ROWS + 50):
        ws.cell(row=i + 1, column=1, value=f"row{i}")
    buf = io.BytesIO()
    wb.save(buf)
    rows = _rows_from_xlsx(buf.getvalue())
    assert len(rows) == _MAX_SCAN_ROWS + 1


def test_rows_from_csv_stops_scanning_past_the_cap():
    from app.bulk_import import _MAX_SCAN_ROWS, _rows_from_csv
    lines = [f"row{i}" for i in range(_MAX_SCAN_ROWS + 50)]
    data = ("\r\n".join(lines) + "\r\n").encode("utf-8")
    rows = _rows_from_csv(data)
    assert len(rows) == _MAX_SCAN_ROWS + 1


def test_import_rejects_oversized_file(admin_client):
    # The route pre-checks size while streaming the upload (_read_upload_with_limit),
    # so run_import()'s own guard is only reachable by calling it directly —
    # exercised here as the defensive check it's meant to be for any other caller.
    from sqlalchemy import select

    from app.bulk_import import MAX_IMPORT_FILE_BYTES, BulkImportError, run_import
    from app.database import SessionLocal
    from app.models import User

    db = SessionLocal()
    try:
        actor = db.scalar(select(User).where(User.email == BOOTSTRAP_EMAIL))
        oversized = b"x" * (MAX_IMPORT_FILE_BYTES + 1)
        try:
            run_import(db, actor, "bugs.xlsx", oversized)
            raise AssertionError("expected BulkImportError")
        except BulkImportError as exc:
            assert "too large" in str(exc).lower()
    finally:
        db.close()


def test_import_rejects_too_many_rows(admin_client):
    from app.bulk_import import MAX_IMPORT_ROWS
    proj = _make_project(admin_client, "Import Proj TooMany")
    rows = [{"Title": f"row {i}", "Project": proj["name"]} for i in range(MAX_IMPORT_ROWS + 1)]
    data = _xlsx_bytes(rows)
    res = _upload(admin_client, data, "bugs.xlsx")
    assert res.status_code == 400
    assert "limit per import" in res.json()["detail"]


def test_create_validated_rows_integrity_error_rolls_back(admin_client, monkeypatch):
    from sqlalchemy import select
    from sqlalchemy.exc import IntegrityError

    from app.bulk_import import BulkImportError, _create_validated_rows, _RowPlan
    from app.database import SessionLocal
    from app.models import Project, User

    proj_json = _make_project(admin_client, "Import Proj IntegrityError")
    db = SessionLocal()
    try:
        actor = db.scalar(select(User).where(User.email == BOOTSTRAP_EMAIL))
        project = db.get(Project, proj_json["id"])
        plan = _RowPlan(
            row_num=3, title="will fail to flush", item_type="Bug",
            project=project, event=None, status="New", priority="Medium",
            environment="DEV", due_date=None, description="", reporter=actor,
        )

        def _boom(*_a, **_k):
            raise IntegrityError("stmt", {}, Exception("simulated"))

        monkeypatch.setattr(db, "flush", _boom)
        try:
            _create_validated_rows(db, actor, [plan])
            raise AssertionError("expected BulkImportError")
        except BulkImportError as exc:
            assert "try again" in str(exc)
    finally:
        db.close()


# --- Project scoping for managers ---

def test_manager_cannot_import_into_project_outside_scope(admin_client):
    outside = _make_project(admin_client, "Admin Only Proj")
    admin_client.post("/api/users", json={
        "name": "Mona Manager", "email": "mona@test.local",
        "role": "manager", "password": "Mana123456",
    })
    # A second, independent TestClient: admin_client's cookie jar can only
    # hold one session at a time, and re-logging-in on it here would silently
    # clobber the admin session the rest of this test still needs.
    from fastapi.testclient import TestClient

    from app.main import app
    manager = TestClient(app)
    res = manager.post("/api/auth/login", json={
        "email": "mona@test.local", "password": "Mana123456",
    })
    assert res.status_code == 200

    data = _xlsx_bytes([{"Title": "sneaky row", "Project": outside["name"]}])
    res = _upload(manager, data, "bugs.xlsx")
    body = res.json()
    assert body["created"] == 0
    assert body["failed"] == 1
    assert "not found" in body["errors"][0]["error"]
