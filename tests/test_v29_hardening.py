"""Regression tests for the v2.9.x security-hardening + perf pass.

Covers:
  1. Webhook SSRF — parse-based host validation rejects the bypasses the
     old substring blocklist missed (IPv6 loopback, 172.16/12, decimal
     IPs, userinfo tricks, *.internal hostnames).
  2. Chatbot Excel export — formula triggers are defanged the same way
     the CSV and reports exports already were.
  3. Attachment downloads — inline rendering is safelist-based; unknown
     MIME types are forced to Content-Disposition: attachment.
  4. Security headers — COOP / CORP / X-Permitted-Cross-Domain-Policies
     present on every response.
  5. Invitation audit IP — X-Forwarded-For only honoured when
     TRUST_PROXY_FORWARDED_FOR is on.
  6. Cached HTML serving still renders both placeholders.
"""
from __future__ import annotations

import pytest


# ---------------------------------------------------------------------------
# 1. Webhook SSRF validation
# ---------------------------------------------------------------------------
class TestWebhookSsrfValidation:
    BLOCKED = [
        "http://localhost/hook",
        "http://127.0.0.1/hook",
        "http://127.1.2.3/hook",
        "http://10.0.0.5/hook",
        "http://192.168.1.1/hook",
        "http://169.254.169.254/latest/meta-data/",
        # New: bypasses the old substring blocklist let through.
        "http://[::1]/hook",                      # IPv6 loopback
        "http://172.16.0.10/hook",                # 172.16/12 private
        "http://172.31.255.254/hook",
        "http://2130706433/hook",                 # decimal 127.0.0.1
        "http://attacker@127.0.0.1/hook",         # userinfo trick
        "http://metadata.google.internal/computeMetadata/v1/",
        "http://foo.internal/hook",
        "http://printer.local/hook",
    ]

    ALLOWED = [
        "https://hooks.example.com/bh",
        "https://example.com:8443/webhook?x=1",
        # Note: documentation ranges (203.0.113.0/24 etc.) are flagged
        # is_private by the ipaddress module and correctly rejected, so
        # a real-world public IP is used here.
        "http://8.8.8.8/hook",
    ]

    @pytest.mark.parametrize("url", BLOCKED)
    def test_blocked_urls_rejected(self, admin_client, url):
        r = admin_client.post("/api/webhooks", json={"name": "h", "url": url})
        assert r.status_code == 422, f"{url} should be rejected, got {r.status_code}"

    @pytest.mark.parametrize("url", ALLOWED)
    def test_public_urls_accepted(self, admin_client, url):
        r = admin_client.post("/api/webhooks", json={"name": "h", "url": url})
        assert r.status_code == 201, f"{url} should be accepted: {r.text}"

    def test_update_url_also_validated(self, admin_client):
        r = admin_client.post("/api/webhooks", json={
            "name": "h", "url": "https://hooks.example.com/bh",
        })
        hook_id = r.json()["id"]
        r = admin_client.put(f"/api/webhooks/{hook_id}", json={
            "url": "http://[::1]/hook",
        })
        assert r.status_code == 422


# ---------------------------------------------------------------------------
# 2. Chatbot Excel formula defang
# ---------------------------------------------------------------------------
class TestChatbotExcelDefang:
    def test_defang_helper(self, app_env):
        from app.chatbot.excel import _defang_formula_text
        assert _defang_formula_text("=cmd|'/c calc'!A1") == "'=cmd|'/c calc'!A1"
        assert _defang_formula_text("+1") == "'+1"
        assert _defang_formula_text("-1") == "'-1"
        assert _defang_formula_text("@x") == "'@x"
        assert _defang_formula_text("\tx") == "'\tx"
        assert _defang_formula_text("normal title") == "normal title"
        assert _defang_formula_text("") == ""

    def test_workbook_cells_defanged(self, app_env):
        import io
        from openpyxl import load_workbook
        from app.chatbot.excel import _build_workbook
        rows = [{
            "id": 1, "title": "=HYPERLINK(\"http://evil\")", "project": "P",
            "status": "New", "priority": "High", "environment": "DEV",
            "reporter": "@victim", "assignees": "+someone", "due_date": "",
            "created_at": "2026-01-01", "updated_at": "2026-01-01",
        }]
        data = _build_workbook(rows, "test export")
        ws = load_workbook(io.BytesIO(data)).active
        # Row 3 is the first data row (banner + header above it).
        assert ws.cell(row=3, column=2).value.startswith("'=")
        assert ws.cell(row=3, column=7).value.startswith("'@")
        assert ws.cell(row=3, column=8).value.startswith("'+")


# ---------------------------------------------------------------------------
# 3. Attachment inline-disposition safelist
# ---------------------------------------------------------------------------
class TestAttachmentDispositionSafelist:
    def _make_bug_with_attachment(self, client, filename, content, ctype):
        r = client.post("/api/projects", json={"name": "P1", "key": "P1"})
        project = r.json()
        r = client.post("/api/bugs", json={
            "project_id": project["id"], "title": "att bug",
            "description": "x", "priority": "Low", "environment": "DEV",
        })
        bug = r.json()
        r = client.post(
            f"/api/bugs/{bug['id']}/attachments",
            files={"file": (filename, content, ctype)},
        )
        assert r.status_code == 201, r.text
        return bug["id"], r.json()["id"]

    def test_image_still_inline(self, admin_client):
        png = (b"\x89PNG\r\n\x1a\n" + b"\x00" * 24)
        bug_id, att_id = self._make_bug_with_attachment(
            admin_client, "x.png", png, "image/png")
        r = admin_client.get(f"/api/bugs/{bug_id}/attachments/{att_id}/download")
        assert r.headers["content-disposition"].startswith("inline;")

    def test_text_plain_still_inline(self, admin_client):
        bug_id, att_id = self._make_bug_with_attachment(
            admin_client, "log.txt", b"hello", "text/plain")
        r = admin_client.get(f"/api/bugs/{bug_id}/attachments/{att_id}/download")
        assert r.headers["content-disposition"].startswith("inline;")

    def test_unknown_type_forced_to_attachment(self, admin_client):
        # application/x-anything is not on the safelist — must download,
        # never render inline.
        bug_id, att_id = self._make_bug_with_attachment(
            admin_client, "blob.bin", b"\x00\x01", "application/x-custom")
        r = admin_client.get(f"/api/bugs/{bug_id}/attachments/{att_id}/download")
        assert r.headers["content-disposition"].startswith("attachment;")

    def test_active_type_still_neutralised(self, admin_client):
        bug_id, att_id = self._make_bug_with_attachment(
            admin_client, "evil.svg", b"<svg/>", "image/svg+xml")
        r = admin_client.get(f"/api/bugs/{bug_id}/attachments/{att_id}/download")
        assert r.headers["content-disposition"].startswith("attachment;")
        assert r.headers["content-type"].startswith("application/octet-stream")


# ---------------------------------------------------------------------------
# 4. Security headers
# ---------------------------------------------------------------------------
class TestSecurityHeaders:
    def test_coop_corp_headers_on_api(self, client):
        r = client.get("/api/health")
        assert r.headers.get("cross-origin-opener-policy") == "same-origin"
        assert r.headers.get("cross-origin-resource-policy") == "same-origin"
        assert r.headers.get("x-permitted-cross-domain-policies") == "none"

    def test_existing_headers_still_present(self, client):
        r = client.get("/api/health")
        assert r.headers.get("x-content-type-options") == "nosniff"
        assert r.headers.get("x-frame-options") == "DENY"
        assert "content-security-policy" in r.headers


# ---------------------------------------------------------------------------
# 5. Invitation audit IP — proxy header not trusted by default
# ---------------------------------------------------------------------------
class TestInvitationClientIp:
    def test_xff_ignored_without_trust_flag(self, app_env):
        from unittest.mock import Mock
        from app.routes.invitations import _client_ip
        req = Mock()
        req.headers = {"x-forwarded-for": "6.6.6.6"}
        req.client = Mock(host="10.1.2.3")
        assert _client_ip(req) == "10.1.2.3"

    def test_xff_honoured_with_trust_flag(self, app_env, monkeypatch):
        from unittest.mock import Mock
        # Settings attributes are read from the environment at class
        # definition, so patch the cached instance rather than the env.
        from app.config import get_settings
        monkeypatch.setattr(
            get_settings(), "TRUST_PROXY_FORWARDED_FOR", True)
        from app.routes.invitations import _client_ip
        req = Mock()
        req.headers = {"x-forwarded-for": "6.6.6.6, 10.0.0.1"}
        req.client = Mock(host="10.1.2.3")
        assert _client_ip(req) == "6.6.6.6"


# ---------------------------------------------------------------------------
# 6. Cached HTML serving
# ---------------------------------------------------------------------------
class TestCachedHtml:
    def test_login_page_renders_placeholders_repeatedly(self, client):
        # Two requests — second is served from the render cache and must
        # be byte-identical with no leaked placeholders.
        r1 = client.get("/login.html")
        r2 = client.get("/login.html")
        assert r1.status_code == r2.status_code == 200
        assert r1.text == r2.text
        assert "__APP_VERSION__" not in r1.text
        assert "__ASSET_VERSION__" not in r1.text

    def test_sw_js_served_from_cache(self, client):
        r1 = client.get("/sw.js")
        r2 = client.get("/sw.js")
        assert r1.status_code == r2.status_code == 200
        assert r1.text == r2.text
        assert r1.headers["service-worker-allowed"] == "/"
