"""Tests for the Reports feature (enterprise port of v2.9).

Enterprise differs from the internal Bug Hunter:
  - No `bootstrap admin` — every test uses signup to create an org admin.
  - All queries are scoped by `actor.org_id`; cross-org isolation is the
    most important invariant to verify.
  - User creation goes through the invitations flow, not direct POST.

The full coverage matrix from the internal `test_reports.py` lives there;
this file ships the smoke + critical-path tests that map cleanly onto
enterprise's signup-based fixtures.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook


_TEST_SECRET = ("TestPass1!",)   # noqa: S105 — hermetic test password


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _make_project(client, name="ReportProj"):
    r = client.post("/api/projects", json={"name": name, "color": "#000000"})
    assert r.status_code == 201, r.text
    return r.json()


def _make_item(client, project_id, **extra):
    body = {
        "title": "reports smoke item",
        "project_id": project_id,
        "priority": "High",
        "environment": "DEV",
        "item_type": "Bug",
    }
    body.update(extra)
    r = client.post("/api/bugs", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _change_status(client, bug_id, new_status):
    r = client.put(f"/api/bugs/{bug_id}", json={"status": new_status})
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------------------
# Auth gate
# ---------------------------------------------------------------------------
def test_reports_unauthenticated_401(client):
    """No session → 401 on every reports endpoint."""
    r = client.get("/api/reports/types")
    assert r.status_code == 401
    r = client.post("/api/reports/run", json={"report_key": "item_detail"})
    assert r.status_code == 401


def test_admin_can_list_report_types(admin_client):
    r = admin_client.get("/api/reports/types")
    assert r.status_code == 200
    body = r.json()
    keys = {t["key"] for t in body["types"]}
    expected = {
        "item_detail", "throughput", "pending_snapshot",
        "status_distribution", "priority_distribution",
        "project_breakdown", "aging", "timeline", "time_to_resolution",
    }
    assert expected.issubset(keys)


# ---------------------------------------------------------------------------
# Item Detail
# ---------------------------------------------------------------------------
def test_item_detail_returns_admin_org_items(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="alpha")
    _make_item(admin_client, p["id"], title="beta")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "item_detail", "filters": {},
    })
    assert r.status_code == 200
    body = r.json()
    titles = {row["title"] for row in body["rows"]}
    assert {"alpha", "beta"}.issubset(titles)


# ---------------------------------------------------------------------------
# Throughput from activity_log
# ---------------------------------------------------------------------------
def test_throughput_counts_resolution_events(admin_client):
    p = _make_project(admin_client)
    b = _make_item(admin_client, p["id"], title="resolvable-one")
    _change_status(admin_client, b["id"], "Resolved")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "throughput", "filters": {},
    })
    assert r.status_code == 200
    body = r.json()
    assert body["summary"]["total_resolved"] >= 1


# ---------------------------------------------------------------------------
# Cross-org isolation — the critical enterprise invariant
# ---------------------------------------------------------------------------
def test_reports_isolate_across_orgs(two_orgs):
    """Org B's admin must NOT see Org A's items in their Reports run."""
    c_a, c_b, _me_a, _me_b = two_orgs
    # Create a project + bug in Org A.
    p_a = c_a.post("/api/projects", json={"name": "Secret A", "color": "#000000"}).json()
    c_a.post("/api/bugs", json={
        "project_id": p_a["id"], "title": "ORG A SECRET BUG",
        "priority": "High", "environment": "DEV", "item_type": "Bug",
    })
    # Org B runs a report — should see ZERO items.
    r = c_b.post("/api/reports/run", json={"report_key": "item_detail", "filters": {}})
    assert r.status_code == 200
    body = r.json()
    titles = {row["title"] for row in body["rows"]}
    assert "ORG A SECRET BUG" not in titles
    assert body["total"] == 0


# ---------------------------------------------------------------------------
# Pending snapshot — currently-open items only (org-scoped)
# ---------------------------------------------------------------------------
def test_pending_snapshot_includes_only_open_items(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="still-open-1")
    _make_item(admin_client, p["id"], title="still-open-2")
    closed = _make_item(admin_client, p["id"], title="will-close")
    _change_status(admin_client, closed["id"], "Closed")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "pending_snapshot", "filters": {},
    })
    assert r.status_code == 200
    titles = {row["title"] for row in r.json()["rows"]}
    assert {"still-open-1", "still-open-2"} == titles


# ---------------------------------------------------------------------------
# Distributions — group by status / priority
# ---------------------------------------------------------------------------
def test_status_distribution_groups_by_status(admin_client):
    p = _make_project(admin_client)
    b1 = _make_item(admin_client, p["id"], title="sd-1")
    _make_item(admin_client, p["id"], title="sd-2")
    _change_status(admin_client, b1["id"], "Resolved")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "status_distribution", "filters": {},
    })
    assert r.status_code == 200
    by_status = {row["status"]: row["count"] for row in r.json()["rows"]}
    assert by_status.get("Resolved") == 1
    assert by_status.get("New") == 1


def test_priority_distribution_groups_by_priority(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="pd-hi", priority="High")
    _make_item(admin_client, p["id"], title="pd-lo", priority="Low")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "priority_distribution", "filters": {},
    })
    assert r.status_code == 200
    by_priority = {row["priority"]: row["count"] for row in r.json()["rows"]}
    assert by_priority.get("High") == 1
    assert by_priority.get("Low") == 1


# ---------------------------------------------------------------------------
# Aging — open items with age buckets
# ---------------------------------------------------------------------------
def test_aging_lists_open_items_with_buckets(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="aging-1")
    _make_item(admin_client, p["id"], title="aging-2")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "aging", "filters": {},
    })
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2
    for row in body["rows"]:
        assert "age_bucket" in row
    assert "by_bucket" in body["summary"]


# ---------------------------------------------------------------------------
# Project breakdown — created / open / resolved per project
# ---------------------------------------------------------------------------
def test_project_breakdown_counts_per_project(admin_client):
    p1 = _make_project(admin_client, name="PB-A")
    p2 = _make_project(admin_client, name="PB-B")
    _make_item(admin_client, p1["id"], title="pb-a-1")
    _make_item(admin_client, p1["id"], title="pb-a-2")
    _make_item(admin_client, p2["id"], title="pb-b-1")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "project_breakdown", "filters": {},
    })
    assert r.status_code == 200
    by_proj = {row["project"]: row for row in r.json()["rows"]}
    assert by_proj["PB-A"]["created"] == 2
    assert by_proj["PB-B"]["created"] == 1


def test_unknown_report_key_returns_400(admin_client):
    r = admin_client.post("/api/reports/run", json={
        "report_key": "not_a_real_report", "filters": {},
    })
    assert r.status_code == 400


# ---------------------------------------------------------------------------
# Timeline — created vs resolved per day (org-scoped)
# ---------------------------------------------------------------------------
def test_timeline_counts_created_and_resolved(admin_client):
    p = _make_project(admin_client)
    a = _make_item(admin_client, p["id"], title="tl-1")
    _make_item(admin_client, p["id"], title="tl-2")
    _change_status(admin_client, a["id"], "Resolved")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "timeline", "filters": {},
    })
    assert r.status_code == 200
    body = r.json()
    assert body["report_key"] == "timeline"
    assert body["summary"]["window_days"] == len(body["rows"])
    assert body["summary"]["total_created"] >= 2
    assert body["summary"]["total_resolved"] >= 1
    for row in body["rows"]:
        assert set(row) >= {"date", "created", "resolved", "delta"}


def test_timeline_explicit_window_in_past_is_empty(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="tl-window")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "timeline",
        "filters": {"date_from": "2000-01-01", "date_to": "2000-01-03"},
    })
    body = r.json()
    assert body["summary"]["window_days"] == 3
    assert body["summary"]["total_created"] == 0


# ---------------------------------------------------------------------------
# Time to Resolution — hours from creation to resolution + stats
# ---------------------------------------------------------------------------
def test_time_to_resolution_reports_resolved_items(admin_client):
    p = _make_project(admin_client)
    b1 = _make_item(admin_client, p["id"], title="ttr-1")
    b2 = _make_item(admin_client, p["id"], title="ttr-2")
    _make_item(admin_client, p["id"], title="ttr-open")  # excluded
    _change_status(admin_client, b1["id"], "Resolved")
    _change_status(admin_client, b2["id"], "Closed")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "time_to_resolution", "filters": {},
    })
    assert r.status_code == 200
    body = r.json()
    assert body["summary"]["count"] == 2
    titles = {row["title"] for row in body["rows"]}
    assert titles == {"ttr-1", "ttr-2"}
    for k in ("average_hours", "median_hours", "p95_hours",
              "fastest_hours", "slowest_hours"):
        assert k in body["summary"]


def test_time_to_resolution_empty_when_nothing_resolved(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="never-resolved")
    r = admin_client.post("/api/reports/run", json={
        "report_key": "time_to_resolution", "filters": {},
    })
    body = r.json()
    assert body["total"] == 0
    assert body["summary"]["count"] == 0
    assert body["summary"]["average_hours"] == 0


def test_time_to_resolution_isolates_across_orgs(two_orgs):
    """Org B must never see Org A's resolution durations."""
    c_a, c_b, *_ = two_orgs
    p_a = c_a.post("/api/projects", json={"name": "TTR A", "color": "#000000"}).json()
    b_a = c_a.post("/api/bugs", json={
        "project_id": p_a["id"], "title": "A-RESOLVED",
        "priority": "High", "environment": "DEV", "item_type": "Bug",
    }).json()
    c_a.put(f"/api/bugs/{b_a['id']}", json={"status": "Resolved"})
    r = c_b.post("/api/reports/run", json={
        "report_key": "time_to_resolution", "filters": {},
    })
    assert r.status_code == 200
    assert r.json()["summary"]["count"] == 0


# ---------------------------------------------------------------------------
# Combined attribute + entity filters (exercises each filter branch)
# ---------------------------------------------------------------------------
def test_item_detail_attribute_and_entity_filters(admin_client):
    me = admin_client.admin_me
    p = _make_project(admin_client, name="FilterProj")
    ev = admin_client.post("/api/events", json={"name": "FilterEvent"}).json()
    target = _make_item(
        admin_client, p["id"], title="filter-target",
        priority="Critical", environment="PROD",
        assignee_ids=[me["id"]], event_id=ev["id"],
    )
    _make_item(admin_client, p["id"], title="filter-decoy",
               priority="Low", environment="DEV")
    today = datetime.now(timezone.utc).date().isoformat()
    r = admin_client.post("/api/reports/run", json={
        "report_key": "item_detail",
        "filters": {
            "priorities": ["Critical"],
            "environments": ["PROD"],
            "assignee_ids": [me["id"]],
            "reporter_ids": [me["id"]],
            "event_id": ev["id"],
            "date_from": today,
            "date_to": today,
        },
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 1
    assert body["rows"][0]["title"] == "filter-target"
    assert body["rows"][0]["id"] == target["id"]


# ---------------------------------------------------------------------------
# XLSX export — round-trips and is org-scoped
# ---------------------------------------------------------------------------
def test_xlsx_export_returns_valid_workbook(admin_client):
    p = _make_project(admin_client)
    _make_item(admin_client, p["id"], title="export-target")
    r = admin_client.post("/api/reports/export.xlsx", json={
        "report_key": "item_detail", "filters": {"label": "Smoke"},
    })
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert "Filters Applied" in wb.sheetnames


def test_xlsx_export_isolates_across_orgs(two_orgs):
    c_a, c_b, *_ = two_orgs
    p_a = c_a.post("/api/projects", json={"name": "Iso A", "color": "#000000"}).json()
    c_a.post("/api/bugs", json={
        "project_id": p_a["id"], "title": "ONLY-IN-A",
        "priority": "High", "environment": "DEV", "item_type": "Bug",
    })
    r = c_b.post("/api/reports/export.xlsx", json={"report_key": "item_detail", "filters": {}})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    text = " ".join(
        str(v) for row in wb[wb.sheetnames[0]].iter_rows(values_only=True)
        for v in row if v is not None
    )
    assert "ONLY-IN-A" not in text


# ---------------------------------------------------------------------------
# Sleuth report intent
# ---------------------------------------------------------------------------
def test_sleuth_report_intent_returns_report_for_admin(admin_client):
    p = _make_project(admin_client)
    b = _make_item(admin_client, p["id"], title="for-sleuth")
    _change_status(admin_client, b["id"], "Resolved")
    r = admin_client.post("/api/chat/ask", json={
        "message": "report of who resolved how many bugs",
    })
    assert r.status_code == 200
    body = r.json()
    assert body["intent"] == "report", body


# ---------------------------------------------------------------------------
# Legacy CSV endpoint must be gone
# ---------------------------------------------------------------------------
def test_legacy_csv_endpoint_removed(admin_client):
    r = admin_client.get("/api/bugs/export.csv")
    # 404 (no handler) or 422 (catchall /{bug_id} treats "export.csv" as
    # not-an-int) both mean the legacy CSV is gone.
    assert r.status_code in (404, 422), r.text
