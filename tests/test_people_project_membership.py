"""People attached to project data must be able to see the project.

Assignees, reporters, sprint members, managers, owners and leads are notified about the work
they are attached to; attaching someone outside the project would leak titles and descriptions
to them. Every write path refuses it with a clear 400 (bulk import skips the assignee with a
warning instead, like its other soft assignee problems). Admins see every project and pass.
"""
from __future__ import annotations

import io

from openpyxl import Workbook

from app.bulk_import import TEMPLATE_HEADERS

_PASSWORD = "Member2026Pw"


def _project(client, name: str) -> dict:
    res = client.post("/api/projects", json={"name": name, "color": "#123456"})
    assert res.status_code == 201, res.text
    return res.json()


def _user(client, name: str, role: str = "user", project_ids: list[int] | None = None) -> dict:
    res = client.post("/api/users", json={
        "name": name, "email": f"{name.lower().replace(' ', '.')}@test.local",
        "role": role, "password": _PASSWORD, "project_ids": project_ids or [],
    })
    assert res.status_code == 201, res.text
    return res.json()


def _setup(client):
    home = _project(client, "Home Project")
    other = _project(client, "Other Project")
    member = _user(client, "Member Mia", project_ids=[home["id"]])
    outsider = _user(client, "Outsider Oli", project_ids=[other["id"]])
    manager_out = _user(client, "Outside Manager", role="manager", project_ids=[other["id"]])
    return home, other, member, outsider, manager_out


def _assert_not_member(res, name: str, project: str) -> None:
    assert res.status_code == 400, res.text
    detail = res.json()["detail"]
    assert name in detail and f"project '{project}'" in detail and "Add them to the project" in detail


def test_bug_create_and_update_refuse_outsiders(admin_client):
    home, other, member, outsider, _ = _setup(admin_client)
    res = admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Leaky title", "assignee_ids": [outsider["id"]],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")

    res = admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Wrong reporter", "reporter_id": outsider["id"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")
    assert "another reporter" in res.json()["detail"]

    bug = admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Fine bug", "assignee_ids": [member["id"]],
    }).json()
    res = admin_client.put(f"/api/bugs/{bug['id']}", json={
        "assignee_ids": [member["id"], outsider["id"]], "version": bug["version"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")

    # Moving the item to a project its assignee can't see is refused too.
    res = admin_client.put(f"/api/bugs/{bug['id']}", json={
        "project_id": other["id"], "version": bug["version"],
    })
    _assert_not_member(res, "Member Mia", "Other Project")
    unchanged = admin_client.get(f"/api/bugs/{bug['id']}").json()
    assert unchanged["project_id"] == home["id"]
    assert [a["id"] for a in unchanged["assignees"]] == [member["id"]]


def test_admins_can_always_be_assigned(admin_client):
    home, *_ = _setup(admin_client)
    me = admin_client.get("/api/auth/me").json()
    res = admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Admin assigned", "assignee_ids": [me["id"]],
    })
    assert res.status_code == 201, res.text


def test_no_notification_reaches_an_outsider(admin_client):
    home, _other, _member, outsider, _ = _setup(admin_client)
    admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Secret roadmap item", "assignee_ids": [outsider["id"]],
    })
    admin_client.post("/api/auth/logout")
    assert admin_client.post("/api/auth/login", json={
        "email": outsider["email"], "password": _PASSWORD,
    }).status_code == 200
    items = admin_client.get("/api/notifications").json()
    items = items.get("items", items) if isinstance(items, dict) else items
    assert not any("Secret roadmap item" in (n.get("body", "") + n.get("title", "")) for n in items)


def _agile(client, project_id: int) -> int:
    assert client.post(f"/api/agile/projects/{project_id}/enable", json={}).status_code == 200
    return client.get(f"/api/agile/projects/{project_id}/settings").json()["board_id"]


def test_agile_work_items_and_sprints_refuse_outsiders(admin_client):
    home, _other, member, outsider, _ = _setup(admin_client)
    board_id = _agile(admin_client, home["id"])

    res = admin_client.post("/api/agile/work-items", json={
        "project_id": home["id"], "title": "Story for outsider", "item_type": "Story",
        "assignee_ids": [outsider["id"]],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")

    res = admin_client.post(f"/api/agile/sprints?board_id={board_id}", json={
        "name": "Leaky sprint", "assignee_ids": [outsider["id"]],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")

    sprint = admin_client.post(f"/api/agile/sprints?board_id={board_id}", json={
        "name": "Team sprint", "assignee_ids": [member["id"]],
    })
    assert sprint.status_code == 201, sprint.text
    sprint = sprint.json()
    res = admin_client.put(f"/api/agile/sprints/{sprint['id']}", json={
        "assignee_ids": [member["id"], outsider["id"]], "version": sprint["version"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")

    res = admin_client.put(f"/api/agile/sprints/{sprint['id']}/capacity", json={"entries": [
        {"user_id": outsider["id"], "capacity_value": 5, "capacity_unit": "points"},
    ]})
    _assert_not_member(res, "Outsider Oli", "Home Project")
    res = admin_client.put(f"/api/agile/sprints/{sprint['id']}/capacity", json={"entries": [
        {"user_id": member["id"], "capacity_value": 5, "capacity_unit": "points"},
    ]})
    assert res.status_code == 200, res.text


def test_owners_and_leads_must_be_members(admin_client):
    home, _other, member, outsider, _ = _setup(admin_client)
    _agile(admin_client, home["id"])

    res = admin_client.post("/api/agile/releases", json={
        "project_id": home["id"], "name": "v1", "owner_id": outsider["id"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")
    release = admin_client.post("/api/agile/releases", json={
        "project_id": home["id"], "name": "v1", "owner_id": member["id"],
    })
    assert release.status_code == 201, release.text
    release = release.json()
    res = admin_client.put(f"/api/agile/releases/{release['id']}", json={
        "owner_id": outsider["id"], "version": release["version"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")

    res = admin_client.post("/api/agile/components", json={
        "project_id": home["id"], "name": "API", "lead_id": outsider["id"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")
    res = admin_client.post("/api/agile/components", json={
        "project_id": home["id"], "name": "API", "default_assignee_id": 999999,
    })
    assert res.status_code == 400 and "does not exist" in res.json()["detail"]
    component = admin_client.post("/api/agile/components", json={
        "project_id": home["id"], "name": "API", "lead_id": member["id"],
    })
    assert component.status_code == 201, component.text

    epic = admin_client.post("/api/agile/work-items", json={
        "project_id": home["id"], "title": "Big epic", "item_type": "Epic",
    }).json()
    detail = admin_client.get(f"/api/agile/epics/{epic['id']}").json()
    res = admin_client.put(f"/api/agile/epics/{epic['id']}", json={
        "owner_id": outsider["id"], "version": detail["version"],
    })
    _assert_not_member(res, "Outsider Oli", "Home Project")
    res = admin_client.put(f"/api/agile/epics/{epic['id']}", json={
        "owner_id": member["id"], "version": detail["version"],
    })
    assert res.status_code == 200 and res.json()["owner_id"] == member["id"]


def test_event_managers_must_see_the_event(admin_client):
    home, _other, _member, _outsider, manager_out = _setup(admin_client)
    res = admin_client.post("/api/events", json={
        "name": "Release party", "project_id": home["id"], "manager_ids": [manager_out["id"]],
    })
    _assert_not_member(res, "Outside Manager", "Home Project")

    res = admin_client.post("/api/events", json={
        "name": "Company offsite", "manager_ids": [manager_out["id"]],
    })
    assert res.status_code == 400
    assert "only admins" in res.json()["detail"]

    me = admin_client.get("/api/auth/me").json()
    ok = admin_client.post("/api/events", json={"name": "Admin event", "manager_ids": [me["id"]]})
    assert ok.status_code == 201, ok.text


def _xlsx(rows: list[dict[str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="banner")
    for col, label in enumerate(TEMPLATE_HEADERS, start=1):
        ws.cell(row=2, column=col, value=label)
    for r_idx, row in enumerate(rows, start=3):
        for col, label in enumerate(TEMPLATE_HEADERS, start=1):
            ws.cell(row=r_idx, column=col, value=row.get(label, ""))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_bulk_import_skips_outsider_assignees_and_rejects_outsider_reporters(admin_client):
    home, _other, member, outsider, _ = _setup(admin_client)
    headers = {h.lower(): h for h in TEMPLATE_HEADERS}
    project_col, title_col = headers["project"], headers["title"]
    assignee_col = next(h for h in TEMPLATE_HEADERS if h.lower().startswith("assignee"))
    reporter_col = next(h for h in TEMPLATE_HEADERS if h.lower() == "reporter email")
    data = _xlsx([
        {project_col: home["name"], title_col: "Imported item",
         assignee_col: f"{member['email']}; {outsider['email']}"},
        {project_col: home["name"], title_col: "Wrong reporter", reporter_col: outsider["email"]},
    ])
    res = admin_client.post("/api/bugs/import", files={
        "file": ("items.xlsx", io.BytesIO(data), "application/octet-stream"),
    })
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 1, body
    text = str(body)
    assert "not a member of project 'Home Project'" in text
    items = admin_client.get("/api/bugs?q=Imported").json()["items"]
    assert [a["id"] for a in items[0]["assignees"]] == [member["id"]]


def test_people_of_another_organization_are_never_attachable(admin_client):
    """Tenancy: another org's user id is "unknown" on every write path, and never a member."""
    from app.database import SessionLocal
    from app.models import User
    from app.people import non_members
    from tests.conftest import sign_up

    home, *_ = _setup(admin_client)
    foreign_client, _me = sign_up("Org Two", "owner@two.test")
    fid = foreign_client.post("/api/users", json={
        "name": "Foreign Fay", "email": "fay@two.test", "role": "manager", "password": _PASSWORD,
    }).json()["id"]

    res = admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Cross-org", "assignee_ids": [fid],
    })
    assert res.status_code == 400 and "Unknown user ids" in res.json()["detail"], res.text
    res = admin_client.post("/api/bugs", json={
        "project_id": home["id"], "title": "Cross-org reporter", "reporter_id": fid,
    })
    assert res.status_code == 400, res.text
    res = admin_client.post("/api/events", json={"name": "Cross", "manager_ids": [fid]})
    assert res.status_code == 400, res.text

    enabled = admin_client.post(f"/api/agile/projects/{home['id']}/enable", json={"feature_flags": {}})
    assert enabled.status_code == 200, enabled.text
    res = admin_client.post(f"/api/agile/sprints?board_id={enabled.json()['board_id']}", json={
        "name": "Sprint X", "goal": "g", "assignee_ids": [fid],
    })
    assert res.status_code == 400, res.text
    res = admin_client.post("/api/agile/components", json={
        "project_id": home["id"], "name": "Core", "lead_id": fid,
    })
    assert res.status_code == 400, res.text

    db = SessionLocal()
    try:
        fay = db.get(User, fid)
        assert non_members(db, [fay], home["id"]) == [fay]
    finally:
        db.close()
