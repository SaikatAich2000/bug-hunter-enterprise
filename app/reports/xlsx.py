"""XLSX writer for reports: main report sheet, optional Items drill-down, Filters
Applied sheet. Used by routes/reports.py and chatbot/excel.py."""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:   # pragma: no cover — broken installs only
    OPENPYXL_AVAILABLE = False

from app.config import get_settings
from app.reports.engine import ReportColumn, ReportResult
from app.util import defang_formula


class XlsxBuildError(Exception):
    """Raised when the workbook can't be assembled (openpyxl missing)."""


_HEADER_FILL = "1F2A44"   # Bug Hunter dark accent
_HEADER_FG = "FFFFFF"
_BANNER_FILL = "C9764F"   # Bug Hunter orange
_BANNER_FG = "FFFFFF"
_ZEBRA_FILL = "F2F4F8"


# Characters a worksheet cannot hold (C0 controls except tab/newline/CR); openpyxl raises on
# them, which failed the whole export for one stray character in old data.
_ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _defang_formula_text(s: str) -> str:
    """Drop characters a worksheet can't hold, then defang a leading formula trigger."""
    return defang_formula(_ILLEGAL_XML_CHARS.sub("", s)) if s else s


def _ensure_openpyxl() -> None:
    if not OPENPYXL_AVAILABLE:
        raise XlsxBuildError(
            "openpyxl is not installed on this server. "
            "Add it to requirements.txt and redeploy."
        )


def _coerce(value: Any) -> Any:
    """Convert a row value to something openpyxl accepts. None becomes '',
    datetimes become ISO strings, and all strings are passed through
    _defang_formula_text to block formula injection."""
    if value is None:
        return ""
    if isinstance(value, bool):
        # bool is a subclass of int — keep this branch above the int one.
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        return _defang_formula_text(value)
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return _defang_formula_text(str(value))


def _cell(ws, value: Any, *, font=None, fill=None, alignment=None):
    cell = WriteOnlyCell(ws, value=value)
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    return cell


def _write_table(
    ws,
    columns: list[ReportColumn],
    rows: list[dict[str, Any]],
    *,
    banner: str,
) -> int:
    """Write a banner, a header and the rows to a write-only sheet (rows stream straight to
    the file instead of building a cell grid in memory). Returns the number of rows written."""
    ncols = max(1, len(columns))
    for idx, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(8, col.width)
    # Banner and header stay in view while scrolling.
    ws.freeze_panes = "A3"

    banner_font = Font(bold=True, color=_BANNER_FG, size=12)
    banner_fill = PatternFill("solid", fgColor=_BANNER_FILL)
    left = Alignment(horizontal="left", vertical="center")
    # Defang banner/headers too, even though they're server-controlled today. The banner
    # fill spans the table's width (write-only sheets can't merge cells).
    ws.append([_cell(ws, _defang_formula_text(banner), font=banner_font, fill=banner_fill, alignment=left)]
              + [_cell(ws, None, fill=banner_fill) for _ in range(ncols - 1)])

    header_font = Font(bold=True, color=_HEADER_FG)
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    ws.append([_cell(ws, _defang_formula_text(col.label), font=header_font, fill=header_fill, alignment=left)
               for col in columns])

    # Data cells are plain values: a per-cell style costs more than the value itself, Excel
    # aligns numbers right on its own, and one conditional-format rule draws the zebra rows.
    keys = [col.key for col in columns]
    for row in rows:
        ws.append([_coerce(row.get(key, "")) for key in keys])
    if len(rows) > 1:
        data = f"A3:{get_column_letter(ncols)}{len(rows) + 2}"
        ws.conditional_formatting.add(data, FormulaRule(
            formula=["MOD(ROW(),2)=0"], fill=PatternFill("solid", bgColor=_ZEBRA_FILL),
        ))
    return len(rows)


_FILTER_LABELS = [
    ("label", "Run Label"),
    ("date_from", "Date From"),
    ("date_to", "Date To"),
    ("item_types", "Item Types"),
    ("statuses", "Statuses"),
    ("priorities", "Priorities"),
    ("environments", "Environments"),
    ("project_ids", "Project IDs"),
    ("assignee_ids", "Assignee IDs"),
    ("reporter_ids", "Reporter IDs"),
    ("event_id", "Event ID"),
    ("include_not_a_bug", "Include 'Not a Bug'"),
    ("text_search", "Text Search"),
]


def _format_filter_display(v: Any) -> str:
    """Render a Filters.to_meta() value for the Filters Applied sheet."""
    if v is None or v == "" or v == [] or v == {}:
        return "(not set)"
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return str(v)


def _format_summary_display(v: Any) -> Any:
    """Render a summary entry for the Filters Applied sheet."""
    if isinstance(v, dict):
        return ", ".join(f"{k}: {vv}" for k, vv in v.items())
    return v


def _write_filters_sheet(ws, result: ReportResult) -> None:
    """The Filters Applied sheet: run details, every filter and the summary values."""
    bold = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 50
    ws.append([_cell(ws, f"{get_settings().APP_NAME} — {result.report_label}", font=Font(bold=True, size=12))])
    ws.append([f"Generated: {datetime.now(timezone.utc).isoformat(timespec='seconds')}"])
    ws.append([f"Total rows: {result.total}"])
    ws.append([])
    ws.append([_cell(ws, "Filter", font=bold), _cell(ws, "Value", font=bold)])
    f = result.filters or {}
    for key, label in _FILTER_LABELS:
        # text_search/label are user free-text — defang like data cells.
        ws.append([label, _coerce(_format_filter_display(f.get(key)))])
    if result.summary:
        ws.append([])
        ws.append([_cell(ws, "Summary", font=bold), _cell(ws, "Value", font=bold)])
        for k, v in result.summary.items():
            ws.append([str(k).replace("_", " ").title(), _coerce(_format_summary_display(v))])


def build_workbook_bytes(result: ReportResult) -> bytes:
    """Serialise a ReportResult to xlsx bytes. Three sheets at most."""
    _ensure_openpyxl()
    wb = Workbook(write_only=True)
    main_ws = wb.create_sheet((result.report_label or "Report")[:31] or "Report")
    banner = f"{result.report_label} — {result.total} row{'' if result.total == 1 else 's'}"
    if result.filters and result.filters.get("label"):
        banner = f"{banner} · {result.filters['label']}"
    _write_table(main_ws, result.columns, result.rows, banner=banner)

    if result.detail_rows:
        detail_ws = wb.create_sheet("Items")
        det_banner = f"{result.report_label} — Items ({len(result.detail_rows)} row{'' if len(result.detail_rows) == 1 else 's'})"
        _write_table(detail_ws, result.detail_columns, result.detail_rows, banner=det_banner)

    _write_filters_sheet(wb.create_sheet("Filters Applied"), result)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "XlsxBuildError",
    "build_workbook_bytes",
    "_defang_formula_text",
    "OPENPYXL_AVAILABLE",
]
