"""Bulk-create bugs/requirements/tasks from an uploaded .xlsx or .csv file.

The accepted column layout mirrors the "Item Detail Export" report
(app/reports/engine.py::_detail_columns) so a spreadsheet exported from this
app can be edited and re-uploaded as-is. System/computed columns (ID,
Created, Updated, Resolved, Resolved By, Days Open, Attachments) are
recognized but ignored on import.

Two-pass design: every row is parsed and validated in memory first (no DB
writes), then only the rows that passed validation are inserted in one
transaction. A row with a hard error (bad project, invalid enum, etc.) is
skipped and reported; an unresolved/deactivated assignee is a soft warning
that just drops that one assignee, not the whole row.

Notifications are written in-app only (``background=None``) — a 500-row
import must never trigger 500 immediate emails or web pushes. Assignees and
reporters still see new items via the notification bell and the daily
digest (app/scheduler.py).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional

from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app import notification_service
from app.access import accessible_project_ids, can_access_project, scope_event_query
from app.config import get_settings
from app.models import Activity, Bug, Event, Project, User
from app.schemas import (
    ALLOWED_ENVIRONMENTS,
    ALLOWED_ITEM_TYPES,
    ALLOWED_PRIORITIES,
    ALLOWED_STATUSES,
    LEGACY_ITEM_TYPES,
    MIN_TITLE_LENGTH,
    STATUSES_BY_TYPE,
    normalize_choice,
    rich_text_to_plain,
    sanitize_html,
    statuses_for_type,
)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    _OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover - broken installs only
    _OPENPYXL_AVAILABLE = False


class BulkImportError(Exception):
    """Raised for a whole-file problem (bad format, too big, no header row).
    Per-row problems are collected instead of raised — see run_import()."""


def _ensure_openpyxl() -> None:
    if not _OPENPYXL_AVAILABLE:
        raise BulkImportError(
            "openpyxl is not installed on this server. Add it to "
            "requirements.txt and redeploy, or upload a .csv file instead."
        )


# Hard ceilings: bound memory/time before we know whether the file is even
# well-formed. 1000 rows comfortably covers the "500 bugs" use case with
# headroom; MAX_IMPORT_FILE_BYTES keeps a garbage/huge upload from being
# fully buffered.
MAX_IMPORT_ROWS = 1000
MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024
_MAX_SCAN_ROWS = MAX_IMPORT_ROWS + 20
_MAX_REPORTED = 200
# bugs.title is VARCHAR(200); PostgreSQL rejects anything longer.
MAX_TITLE_LENGTH = 200

_BANNER_FILL = "C9764F"   # Bug Hunter orange
_BANNER_FG = "FFFFFF"
_HEADER_FILL = "1F2A44"   # Bug Hunter dark accent
_HEADER_FG = "FFFFFF"

# Exact column order/labels of the Item Detail Export report, so a report
# downloaded from this app can be filled in and re-uploaded unchanged.
TEMPLATE_HEADERS = [
    "ID", "Type", "Title", "Project", "Event", "Status", "Priority", "Env",
    "Reporter", "Reporter Email", "Assignees", "Due Date", "Created",
    "Updated", "Resolved", "Resolved By", "Days Open", "Attachments",
    "Description",
]

# Header label (lowercased) -> Bug field name; None = recognized but not used
# for creation (system/computed columns). Keeping them in this map means a
# previously exported report re-uploaded as-is doesn't trip an "unknown
# column" complaint — those columns are just silently skipped.
_LABEL_TO_FIELD: dict[str, Optional[str]] = {
    "id": None,
    "type": "item_type",
    "title": "title",
    "project": "project",
    "event": "event",
    "status": "status",
    "priority": "priority",
    "env": "environment",
    "reporter": "reporter",
    "reporter email": "reporter_email",
    "assignees": "assignees",
    "due date": "due_date",
    "created": None,
    "updated": None,
    "resolved": None,
    "resolved by": None,
    "days open": None,
    "attachments": None,
    "description": "description",
}

# The header row must contain at least these to be recognized as such.
_REQUIRED_HEADER_TOKENS = {"title", "project"}


def _stringify_cell(value: Any) -> str:
    """Cell value as trimmed text; a date/datetime cell (Excel auto-formats
    typed dates) becomes YYYY-MM-DD instead of Python's repr."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    _ensure_openpyxl()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is a clean 400
        raise BulkImportError("Could not read this file as an Excel workbook.") from exc
    ws = wb.worksheets[0]
    rows: list[list[str]] = []
    for raw_row in ws.iter_rows(values_only=True):
        rows.append([_stringify_cell(v) for v in raw_row])
        if len(rows) > _MAX_SCAN_ROWS:
            break
    return rows


def _rows_from_csv(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BulkImportError("This CSV file isn't valid UTF-8 text.") from exc
    rows: list[list[str]] = []
    for raw_row in csv.reader(io.StringIO(text)):
        rows.append([_stringify_cell(v) for v in raw_row])
        if len(rows) > _MAX_SCAN_ROWS:
            break
    return rows


def _find_header(rows: list[list[str]]) -> tuple[int, dict[int, str]]:
    """Locate the header row (banner rows before it are skipped) and map
    column index -> field name. Raises if no row has Title + Project."""
    for i, row in enumerate(rows[:5]):
        col_map: dict[int, str] = {}
        found_labels: set[str] = set()
        for col_idx, raw in enumerate(row):
            label = raw.strip().lower()
            if label not in _LABEL_TO_FIELD:
                continue
            found_labels.add(label)
            field_name = _LABEL_TO_FIELD[label]
            if field_name:
                col_map[col_idx] = field_name
        if _REQUIRED_HEADER_TOKENS.issubset(found_labels):
            return i, col_map
    raise BulkImportError(
        "Could not find a header row with the expected columns (needs at "
        "least 'Title' and 'Project'). Download the template and use its "
        "exact column headers."
    )


@dataclass
class _Lookups:
    projects_by_name: dict[str, Project]
    events_by_name: dict[str, list[Event]]
    users_by_email: dict[str, User]
    users_by_name: dict[str, list[User]]


def _load_lookups(db: Session, actor: User, accessible: set[int]) -> _Lookups:
    """One bulk read per entity type instead of a query per row. Everything is limited to
    what the importer may see in their own organization: an email or name in the file can
    never resolve to a user, project or event of another tenant."""
    proj_stmt = select(Project).where(Project.id.in_(accessible))
    projects_by_name = {p.name.lower(): p for p in db.scalars(proj_stmt).all()}

    ev_stmt = scope_event_query(select(Event), actor, accessible)
    events_by_name: dict[str, list[Event]] = {}
    for e in db.scalars(ev_stmt).all():
        events_by_name.setdefault(e.name.lower(), []).append(e)

    users = list(db.scalars(select(User).where(User.org_id == actor.org_id)).all())
    users_by_email = {u.email.lower(): u for u in users}
    users_by_name: dict[str, list[User]] = {}
    for u in users:
        users_by_name.setdefault(u.name.lower(), []).append(u)

    return _Lookups(projects_by_name, events_by_name, users_by_email, users_by_name)


@dataclass
class _RowPlan:
    row_num: int
    title: str
    item_type: str
    project: Project
    event: Optional[Event]
    status: str
    priority: str
    environment: str
    due_date: Optional[str]
    description: str
    reporter: User
    assignees: list[User] = field(default_factory=list)
    warning: str = ""


def _resolve_row_title(cells: dict[str, str], row_num: int, errors: list[str]) -> str:
    title = cells.get("title", "").strip()
    if not title:
        # A blank title never fails the row — it's just not identifying.
        return f"Untitled import — row {row_num}"
    if len(title) < MIN_TITLE_LENGTH:
        errors.append(f"Title must be at least {MIN_TITLE_LENGTH} characters")
    elif len(title) > MAX_TITLE_LENGTH:
        errors.append(f"Title must be at most {MAX_TITLE_LENGTH} characters")
    return title


def _resolve_row_item_type(cells: dict[str, str], errors: list[str]) -> str:
    raw_type = cells.get("item_type", "").strip() or "Bug"
    try:
        item_type = normalize_choice(raw_type, ALLOWED_ITEM_TYPES, "Type")
        if item_type not in LEGACY_ITEM_TYPES:
            raise ValueError(
                f"Invalid Type '{raw_type}'. Bulk import only supports: "
                f"{', '.join(LEGACY_ITEM_TYPES)} (Epic/Story/Sub-task are "
                f"created via the Agile module)"
            )
    except ValueError as exc:
        item_type = "Bug"
        errors.append(str(exc))
    return item_type


def _resolve_row_project(
    cells: dict[str, str], lookups: _Lookups, accessible: set[int], errors: list[str],
) -> Optional[Project]:
    project_name = cells.get("project", "").strip()
    if project_name:
        project = lookups.projects_by_name.get(project_name.lower())
        if project is None or not can_access_project(accessible, project.id):
            errors.append(f"Project '{project_name}' not found")
        return project
    # No project given: fine as long as there's exactly one to default to
    # (already scoped to the importer's accessible projects) — otherwise
    # there's no safe guess, so the row is reported instead of misfiled.
    candidates = list(lookups.projects_by_name.values())
    if len(candidates) == 1:
        return candidates[0]
    errors.append(
        "Project is required (leave it blank only when there is "
        "exactly one project to choose from)"
    )
    return None


def _resolve_row_event(
    cells: dict[str, str], lookups: _Lookups, errors: list[str],
) -> Optional[Event]:
    event_name = cells.get("event", "").strip()
    if not event_name:
        return None
    matches = lookups.events_by_name.get(event_name.lower(), [])
    if not matches:
        errors.append(f"Event '{event_name}' not found")
        return None
    if len(matches) > 1:
        errors.append(
            f"Event name '{event_name}' matches more than one event — "
            "rename one or leave this blank"
        )
        return None
    return matches[0]


def _resolve_row_status(cells: dict[str, str], item_type: str, errors: list[str]) -> str:
    raw_status = cells.get("status", "").strip()
    status_val = "New"
    if raw_status:
        try:
            status_val = normalize_choice(raw_status, ALLOWED_STATUSES, "Status")
        except ValueError as exc:
            errors.append(str(exc))
    allowed_for_type = statuses_for_type(item_type)
    if status_val not in allowed_for_type:
        errors.append(
            f"Status '{status_val}' is not valid for {item_type}. "
            f"Allowed: {', '.join(allowed_for_type)}"
        )
    return status_val


def _resolve_row_priority(cells: dict[str, str], errors: list[str]) -> str:
    raw_priority = cells.get("priority", "").strip()
    priority = "Medium"
    if raw_priority:
        try:
            priority = normalize_choice(raw_priority, ALLOWED_PRIORITIES, "Priority")
        except ValueError as exc:
            errors.append(str(exc))
    return priority


def _resolve_row_environment(cells: dict[str, str], errors: list[str]) -> str:
    raw_env = cells.get("environment", "").strip()
    environment = "DEV"
    if raw_env:
        try:
            environment = normalize_choice(raw_env, ALLOWED_ENVIRONMENTS, "Env")
        except ValueError as exc:
            errors.append(str(exc))
    return environment


def _resolve_row_due_date(cells: dict[str, str], errors: list[str]) -> Optional[str]:
    due_raw = cells.get("due_date", "").strip()
    if not due_raw:
        return None
    try:
        datetime.strptime(due_raw, "%Y-%m-%d")
        return due_raw
    except ValueError:
        errors.append("Due Date must be YYYY-MM-DD")
        return None


def _find_reporter_by_email(
    reporter_email: str, lookups: _Lookups, errors: list[str],
) -> Optional[User]:
    """Look up a reporter by email address."""
    if not reporter_email:
        return None
    reporter = lookups.users_by_email.get(reporter_email)
    if reporter is None:
        errors.append(f"Reporter email '{reporter_email}' not found")
    return reporter


def _find_reporter_by_name(
    reporter_name: str, lookups: _Lookups, errors: list[str],
) -> Optional[User]:
    """Look up a reporter by name (with multi-match detection)."""
    if not reporter_name:
        return None
    matches = lookups.users_by_name.get(reporter_name.lower(), [])
    if not matches:
        errors.append(f"Reporter '{reporter_name}' not found")
        return None
    if len(matches) > 1:
        errors.append(
            f"Reporter name '{reporter_name}' matches more than one "
            "user — use Reporter Email instead"
        )
        return None
    return matches[0]


def _resolve_row_reporter(
    cells: dict[str, str], lookups: _Lookups, actor: User, errors: list[str],
) -> Optional[User]:
    """Resolve a reporter by email or name, defaulting to the actor."""
    reporter_email = cells.get("reporter_email", "").strip().lower()
    reporter_name = cells.get("reporter", "").strip()
    
    # Try email lookup first (most specific)
    reporter = _find_reporter_by_email(reporter_email, lookups, errors)
    if reporter is not None:
        # Check if active
        if not reporter.is_active:
            errors.append(f"Inactive Reporter '{reporter.name}' — deactivated")
            return None
        return reporter
    
    # Try name lookup second
    reporter = _find_reporter_by_name(reporter_name, lookups, errors)
    if reporter is not None:
        # Check if active
        if not reporter.is_active:
            errors.append(f"Inactive Reporter '{reporter.name}' — deactivated")
            return None
        return reporter
    
    # Default to the actor (current user) if no explicit reporter
    if not reporter_email and not reporter_name:
        return actor
    
    # Error path: explicit reporter specified but not found/active
    return None


def _resolve_one_assignee_token(
    token: str, lookups: _Lookups, warn_parts: list[str],
) -> Optional[User]:
    user: Optional[User] = None
    if "@" in token:
        user = lookups.users_by_email.get(token.lower())
    else:
        name_matches = lookups.users_by_name.get(token.lower(), [])
        if len(name_matches) == 1:
            user = name_matches[0]
        elif len(name_matches) > 1:
            warn_parts.append(
                f"'{token}' matches more than one user — skipped, "
                "use an email instead"
            )
            return None
    if user is None:
        warn_parts.append(f"assignee '{token}' not found — skipped")
        return None
    if not user.is_active:
        warn_parts.append(f"assignee '{token}' is deactivated — skipped")
        return None
    return user


def _resolve_row_assignees(
    cells: dict[str, str], lookups: _Lookups,
) -> tuple[list[User], list[str]]:
    # Assignees are forgiving: an unresolved/deactivated name only drops that
    # one assignee (warning), it never fails the whole row.
    assignees: list[User] = []
    warn_parts: list[str] = []
    raw_assignees = cells.get("assignees", "").strip()
    if raw_assignees:
        for token in re.split(r"[;,]", raw_assignees):
            token = token.strip()
            if not token:
                continue
            user = _resolve_one_assignee_token(token, lookups, warn_parts)
            if user is not None and user not in assignees:
                assignees.append(user)
    return assignees, warn_parts


def _process_row(
    row_num: int,
    cells: dict[str, str],
    lookups: _Lookups,
    accessible: set[int],
    actor: User,
) -> tuple[Optional[_RowPlan], list[str]]:
    """Validate one row. Returns (plan, []) on success or (None, errors).

    Every column is optional — a blank cell falls back to a sensible default
    (or, for Project, the importer's only accessible project if there's
    exactly one). A hard error only fires when a value was actually given
    and it's invalid/unresolvable/ambiguous, never merely for being blank."""
    errors: list[str] = []

    title = _resolve_row_title(cells, row_num, errors)
    item_type = _resolve_row_item_type(cells, errors)
    project = _resolve_row_project(cells, lookups, accessible, errors)
    event = _resolve_row_event(cells, lookups, errors)
    status_val = _resolve_row_status(cells, item_type, errors)
    priority = _resolve_row_priority(cells, errors)
    environment = _resolve_row_environment(cells, errors)
    due_date = _resolve_row_due_date(cells, errors)
    reporter = _resolve_row_reporter(cells, lookups, actor, errors)
    assignees, warn_parts = _resolve_row_assignees(cells, lookups)

    if errors:
        return None, errors
    assert project is not None
    assert reporter is not None

    # Plain text, like every other work-item create path.
    description = rich_text_to_plain(sanitize_html(cells.get("description", "").strip()))
    plan = _RowPlan(
        row_num=row_num,
        title=title, item_type=item_type, project=project, event=event,
        status=status_val, priority=priority, environment=environment,
        due_date=due_date, description=description, reporter=reporter,
        assignees=assignees, warning="; ".join(warn_parts),
    )
    return plan, []


@dataclass
class BulkImportOutcome:
    created: int
    failed: int
    created_ids: list[int]
    errors: list[dict[str, Any]]
    warnings: list[dict[str, Any]]
    errors_truncated: bool
    warnings_truncated: bool
    message: str


def _create_validated_rows(
    db: Session, actor: User, plans: list[_RowPlan],
) -> list[int]:
    """Insert every validated row in one transaction; audit-logged and
    in-app-notified the same way a manual create is, minus the per-item
    email/push (see module docstring)."""
    bug_plan_pairs: list[tuple[Bug, _RowPlan]] = []
    # Rows are added in file order; the flush ranks each at the bottom of its
    # project's backlog in that order (app/agile/integrity.py).
    for plan in plans:
        pid = plan.project.id
        bug = Bug(
            project_id=pid,
            title=plan.title,
            description=plan.description,
            reporter_id=plan.reporter.id,
            item_type=plan.item_type,
            status=plan.status,
            priority=plan.priority,
            environment=plan.environment,
            due_date=plan.due_date,
            event_id=plan.event.id if plan.event else None,
        )
        bug.assignees = list(plan.assignees)
        db.add(bug)
        bug_plan_pairs.append((bug, plan))

    try:
        db.flush()
    except SQLAlchemyError as exc:
        db.rollback()
        raise BulkImportError(
            "Import failed unexpectedly while saving; no items were "
            "created. Please try again."
        ) from exc

    created_ids: list[int] = []
    for bug, plan in bug_plan_pairs:
        created_ids.append(bug.id)
        db.add(Activity(
            org_id=actor.org_id, bug_id=bug.id, entity_type="bug", entity_id=bug.id,
            actor_user_id=actor.id, actor_name=actor.name, action="bug_created",
            detail=(
                f"{bug.item_type} #{bug.id} '{bug.title}' created with "
                f"status '{bug.status}'. (bulk import)"
            ),
        ))
        if plan.assignees:
            names = ", ".join(a.name for a in plan.assignees)
            db.add(Activity(
                org_id=actor.org_id, bug_id=bug.id, entity_type="bug", entity_id=bug.id,
                actor_user_id=actor.id, actor_name=actor.name,
                action="assignees_added",
                detail=f"Bug #{bug.id} '{bug.title}' assigned to: {names}",
            ))
        assignee_ids = [a.id for a in plan.assignees]
        notification_service.notify(
            db, assignee_ids, kind="assigned", background=None,
            title=f"Assigned to {bug.item_type.lower()} #{bug.id}",
            body=f"{actor.name} assigned you to “{bug.title}” (bulk import).",
            bug_id=bug.id, actor_name=actor.name,
        )
        if plan.reporter.id != actor.id and plan.reporter.id not in assignee_ids:
            notification_service.notify(
                db, [plan.reporter.id], kind="reported", background=None,
                title=f"You're the reporter on {bug.item_type.lower()} #{bug.id}",
                body=(
                    f"{actor.name} filed “{bug.title}” with you "
                    "as reporter (bulk import)."
                ),
                bug_id=bug.id, actor_name=actor.name,
            )
    db.commit()
    return created_ids


def _parse_import_rows(filename: str, data: bytes) -> list[list[str]]:
    if len(data) > MAX_IMPORT_FILE_BYTES:
        raise BulkImportError(
            f"File is too large. The limit is {MAX_IMPORT_FILE_BYTES // (1024 * 1024)} MB."
        )
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "xlsx":
        return _rows_from_xlsx(data)
    if ext == "csv":
        return _rows_from_csv(data)
    raise BulkImportError("Unsupported file type. Upload a .xlsx or .csv file.")


def _numbered_data_rows(rows: list[list[str]]) -> tuple[int, dict[int, str], list[tuple[int, list[str]]]]:
    header_idx, col_map = _find_header(rows)
    data_rows = rows[header_idx + 1:]
    # (sheet row number, raw row) for non-blank rows only; blank trailing
    # rows (common in spreadsheet exports) don't count against the limit.
    numbered = [
        (i, r) for i, r in enumerate(data_rows, start=header_idx + 2)
        if any(c.strip() for c in r)
    ]
    if len(numbered) > MAX_IMPORT_ROWS:
        raise BulkImportError(
            f"This file has {len(numbered)} data rows; the limit per import "
            f"is {MAX_IMPORT_ROWS}. Split it into smaller files and upload "
            "separately."
        )
    return header_idx, col_map, numbered


def _validate_all_rows(
    numbered: list[tuple[int, list[str]]],
    col_map: dict[int, str],
    lookups: _Lookups,
    accessible: set[int],
    actor: User,
) -> tuple[list[_RowPlan], list[dict[str, Any]], list[dict[str, Any]]]:
    plans: list[_RowPlan] = []
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for row_num, raw_row in numbered:
        cells = {
            field_name: (raw_row[idx] if idx < len(raw_row) else "")
            for idx, field_name in col_map.items()
        }
        title_hint = cells.get("title", "").strip()[:80]
        plan, row_errors = _process_row(row_num, cells, lookups, accessible, actor)
        if row_errors:
            errors.append({"row": row_num, "title": title_hint, "error": "; ".join(row_errors)})
            continue
        assert plan is not None
        plans.append(plan)
        if plan.warning:
            warnings.append({"row": row_num, "title": plan.title[:80], "warning": plan.warning})
    return plans, errors, warnings


def _import_outcome_message(created: int, failed: int) -> str:
    if created and not failed:
        return f"Imported {created} item{'s' if created != 1 else ''}."
    if created and failed:
        return f"Imported {created} item{'s' if created != 1 else ''}; {failed} row(s) skipped — see details below."
    if not created and not failed:
        return "No data rows found in this file."
    return f"No items were imported; {failed} row(s) had errors."


def run_import(db: Session, actor: User, filename: str, data: bytes) -> BulkImportOutcome:
    """Parse, validate, and create. Raises BulkImportError for a whole-file
    problem; per-row problems are collected in the returned outcome."""
    rows = _parse_import_rows(filename, data)
    _header_idx, col_map, numbered = _numbered_data_rows(rows)

    accessible = accessible_project_ids(db, actor)
    lookups = _load_lookups(db, actor, accessible)

    plans, errors, warnings = _validate_all_rows(numbered, col_map, lookups, accessible, actor)
    created_ids = _create_validated_rows(db, actor, plans) if plans else []

    created, failed = len(created_ids), len(errors)
    message = _import_outcome_message(created, failed)

    return BulkImportOutcome(
        created=created,
        failed=failed,
        created_ids=created_ids,
        errors=errors[:_MAX_REPORTED],
        warnings=warnings[:_MAX_REPORTED],
        errors_truncated=len(errors) > _MAX_REPORTED,
        warnings_truncated=len(warnings) > _MAX_REPORTED,
        message=message,
    )


def _write_instructions(ws) -> None:
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90
    rows: list[tuple[str, str]] = [
        ("Bulk Import — how to fill this template", ""),
        ("", ""),
        (
            "Every column is optional",
            "A blank cell falls back to a sensible default; a row only fails "
            "when a value was actually given and it's invalid or unresolvable.",
        ),
        ("", ""),
        ("Recommended columns", ""),
        ("Title", "At least 3 characters if given; a blank Title gets a placeholder like \"Untitled import — row 7\"."),
        (
            "Project",
            "Must exactly match an existing project name (case-insensitive) if "
            "given. Left blank, it defaults to your only accessible project — "
            "if there's more than one, the row is reported instead of guessed.",
        ),
        ("", ""),
        ("Other optional columns", ""),
        ("Type", f"One of: {', '.join(LEGACY_ITEM_TYPES)}. Default: Bug."),
        ("Status", "Must be valid for the row's Type (see the table below). Default: New."),
        ("Priority", f"One of: {', '.join(ALLOWED_PRIORITIES)}. Default: Medium."),
        ("Env", f"One of: {', '.join(ALLOWED_ENVIRONMENTS)}. Default: DEV."),
        ("Event", "Must exactly match an existing event name; leave blank for none."),
        ("Reporter Email", "An existing user's email. Takes priority over Reporter."),
        ("Reporter", "An existing user's name, used only if Reporter Email is blank. Defaults to you if both are blank."),
        ("Assignees", 'Comma-separated names or emails, e.g. "Jane Doe, john@example.com". An unmatched name is skipped with a warning, not a failure.'),
        ("Due Date", "YYYY-MM-DD."),
        ("Description", "Plain text or simple HTML."),
        ("", ""),
        ("Ignored on import", ""),
        (
            "ID, Created, Updated, Resolved, Resolved By, Days Open, Attachments",
            "System-generated; leave blank. Any value here is not read.",
        ),
        ("", ""),
        ("Valid statuses by Type", ""),
    ]
    for item_type, statuses in STATUSES_BY_TYPE.items():
        if item_type not in LEGACY_ITEM_TYPES:
            continue
        rows.append((item_type, ", ".join(statuses)))
    for r_idx, (a, b) in enumerate(rows, start=1):
        cell_a = ws.cell(row=r_idx, column=1, value=a)
        if a and not b:
            cell_a.font = Font(bold=True)
        ws.cell(row=r_idx, column=2, value=b)


def build_template_workbook() -> bytes:
    """Blank workbook: banner + header row (0 data rows) matching the Item
    Detail Export columns exactly, plus an Instructions sheet."""
    _ensure_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = "Bulk Import Template"
    ncols = len(TEMPLATE_HEADERS)
    banner = f"{get_settings().APP_NAME} — Bulk Import Template ({ncols} columns)"
    banner_cell = ws.cell(row=1, column=1, value=banner)
    banner_cell.font = Font(bold=True, color=_BANNER_FG, size=12)
    banner_cell.fill = PatternFill("solid", fgColor=_BANNER_FILL)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 22

    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    header_font = Font(bold=True, color=_HEADER_FG)
    for idx, label in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=2, column=idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        ws.column_dimensions[get_column_letter(idx)].width = 22
    ws.freeze_panes = ws.cell(row=3, column=1)

    _write_instructions(wb.create_sheet("Instructions"))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "BulkImportError",
    "BulkImportOutcome",
    "MAX_IMPORT_FILE_BYTES",
    "MAX_IMPORT_ROWS",
    "TEMPLATE_HEADERS",
    "build_template_workbook",
    "run_import",
]
